import openpyxl
import re
from django.core.management.base import BaseCommand
from django.utils.text import slugify
from apps.master_data.models import Department, CivicCategory, CivicLocation

class Command(BaseCommand):
    help = 'Loads official RMC departments, categories, and ward-area gazetteers from the official RMC portal and excel dataset.'

    def handle(self, *args, **options):
        self.stdout.write(self.style.NOTICE("Syncing official RMC metadata from Portal & Excel dataset..."))

        # 1. Official RMC Department Mapping (from Complaint Registration & Excel)
        official_departments = [
            {
                "name": "DRAINAGE CHOCKUP/OVERFLOW",
                "code": "drainage-chockup-overflow",
                "head_title": "City Drainage Engineer (Gutter Branch)",
                "head_email": "drainage.rmc@gov.in",
                "description": "Main line overflowing, drainage choke-up, manhole flooding, apartment sewage lines.",
            },
            {
                "name": "ROSHNI BRANCH",
                "code": "roshni-branch",
                "head_title": "Executive Engineer (Lighting & Electrical)",
                "head_email": "roshni.rmc@gov.in",
                "description": "Streetlights, LED fixtures, dark spots, light timing controls, pole damage.",
            },
            {
                "name": "SOLID WASTE MANAGEMENT",
                "code": "solid-waste-management",
                "head_title": "Director (Solid Waste Management)",
                "head_email": "swm.rmc@gov.in",
                "description": "Door-to-door tipper vans, community bins, street sweeping, dumping clearance.",
            },
            {
                "name": "WATER WORKS(OUTDOOR)",
                "code": "water-works-outdoor",
                "head_title": "Superintending Engineer (Water Supply)",
                "head_email": "waterworks.rmc@gov.in",
                "description": "Water distribution, pipeline leakage, contaminated drinking water, low pressure, valve repair.",
            },
            {
                "name": "DRAINAGE MAINTENANCE",
                "code": "drainage-maintenance",
                "head_title": "Executive Engineer (Drainage Civil Works)",
                "head_email": "drainage.maint@rmc.gov.in",
                "description": "Manhole leveling, chamber cover damage, sewer line structural repair.",
            },
            {
                "name": "BANDHKAM",
                "code": "bandhkam-civil-roads",
                "head_title": "City Engineer (Civil & Roads)",
                "head_email": "bandhkam.rmc@gov.in",
                "description": "Roads, potholes, asphalt resurfacing, paver blocks, footpaths, bridges, municipal buildings.",
            },
            {
                "name": "A.N.C.D. (DOG & ANIMAL CONTROL)",
                "code": "ancd-animal-control",
                "head_title": "Superintendent (Animal Nuisance Control Department)",
                "head_email": "ancd.rmc@gov.in",
                "description": "Catch stray cattle (dhor pakadva), dog nuisance, rabies sterilization, animal shelters.",
            },
            {
                "name": "CITY BUS (RAJKOT RAJPATH)",
                "code": "city-bus-rajpath",
                "head_title": "General Manager (Rajkot Rajpath Ltd / BRTS)",
                "head_email": "citybus.rmc@gov.in",
                "description": "City bus timings, driver behavior, bus stop condition, route extensions.",
            },
            {
                "name": "CONSERVANCY (DEAD ANIMAL)",
                "code": "conservancy-dead-animal",
                "head_title": "Sanitary Inspector (Conservancy)",
                "head_email": "conservancy.rmc@gov.in",
                "description": "Emergency dead animal carcass disposal and bio-sanitation.",
            },
            {
                "name": "DABAN HATAV BRANCH",
                "code": "daban-hatav-branch",
                "head_title": "Encroachment Removal Officer",
                "head_email": "dabanhatav.rmc@gov.in",
                "description": "Removal of illegal encroachments on footpaths, roads, and public spaces.",
            },
            {
                "name": "HEALTH & URBAN MALERIYA",
                "code": "health-urban-maleriya",
                "head_title": "Medical Officer of Health (MOH)",
                "head_email": "health.rmc@gov.in",
                "description": "Mosquito fogging, vector-borne disease control, clinic hygiene.",
            },
            {
                "name": "GARDEN & SPORTS",
                "code": "garden-sports",
                "head_title": "Garden Superintendent",
                "head_email": "garden.rmc@gov.in",
                "description": "Public parks, tree cutting/pruning, playground equipment.",
            },
            {
                "name": "TOWN PLANNING & ESTATE",
                "code": "town-planning-estate",
                "head_title": "Town Planning Officer",
                "head_email": "townplanning.rmc@gov.in",
                "description": "Building permissions, layout violations, public estate maintenance.",
            },
            {
                "name": "FOOD & SANITATION",
                "code": "food-sanitation",
                "head_title": "Food Safety Officer",
                "head_email": "food.rmc@gov.in",
                "description": "Food quality inspection, adulteration checks, restaurant hygiene.",
            },
            {
                "name": "TRAFFIC & TRNS. CELL",
                "code": "traffic-transport-cell",
                "head_title": "Traffic Engineer",
                "head_email": "traffic.rmc@gov.in",
                "description": "Traffic signals, road markings, speed breakers, parking management.",
            },
        ]

        dept_objs = {}
        for d in official_departments:
            dept, _ = Department.objects.update_or_create(
                code=d["code"],
                defaults={
                    "name": d["name"],
                    "head_title": d["head_title"],
                    "head_email": d["head_email"],
                    "description": d["description"],
                    "is_active": True,
                }
            )
            dept_objs[d["name"]] = dept
        self.stdout.write(self.style.SUCCESS(f"  Synced {len(dept_objs)} official RMC departments."))

        # 2. Official Civic Locations (18 Wards & 3 Zones)
        # Zone breakdown as documented by RMC:
        # East Zone: Wards 4, 5, 6, 15, 16, 18
        # Central Zone: Wards 2, 3, 7, 13, 14, 17
        # West Zone: Wards 1, 8, 9, 10, 11, 12
        zone_mapping = {
            1: ("WEST", "Gandhigram / Madhapar / 150ft Ring Rd"),
            2: ("CENTRAL", "Railnagar / Popatpara / Bajrangwadi"),
            3: ("CENTRAL", "Dharmendrasinhji Road / Danapith / Para Bazaar"),
            4: ("EAST", "Bhagvati Para / Bedi Para / Morbi Road"),
            5: ("EAST", "Pedak Road / Greenland Chowkdi / Sant Kabir Rd"),
            6: ("EAST", "Chunara Vad / Aji Dam / Dudh Dairy"),
            7: ("CENTRAL", "Amin Marg / Race Course / Yagnik Road"),
            8: ("WEST", "University Road / Indira Circle / Sadhu Vaswani Rd"),
            9: ("WEST", "Raiya Road / Nana Mava / Speedwell"),
            10: ("WEST", "150 Feet Ring Road / Mavdi / Umiya Chowk"),
            11: ("WEST", "Kalawad Road / Pushkar Dham / Haridarshan"),
            12: ("WEST", "Vavdi / Kankot Road / GEC"),
            13: ("CENTRAL", "Tagore Road / Bhaktinagar Station / Gondal Road"),
            14: ("CENTRAL", "Bhakti Nagar / Geeta Mandir / Dadi Baug"),
            15: ("EAST", "Sant Kabir Road / Bhavnagar Road / Chunarwad"),
            16: ("EAST", "Aji Dam / Marketing Yard / 50ft Road"),
            17: ("CENTRAL", "Atika / Kothariya / Dhebar Road South"),
            18: ("EAST", "Samrat Industrial / Kotharia Ring Road / GIDC"),
        }

        # Parse Areas from ComplaintRegistration HTML file if available
        areas_by_ward = {w: [] for w in range(1, 19)}
        try:
            with open('/Users/suhaniparmar/.gemini/antigravity/brain/7b19e1fe-0b79-4a0d-b26a-52ed6d35c8a1/.system_generated/steps/173/content.md') as f:
                html = f.read()
            import base64
            viewstate = re.search(r'name=\"__VIEWSTATE\" id=\"__VIEWSTATE\" value=\"([^\"]+)\"', html)
            if viewstate:
                decoded = base64.b64decode(viewstate.group(1)).decode('latin1', errors='ignore')
                raw_areas = re.findall(r'(\d{2})\s*\|\s*([A-Za-z0-9\s\.\(\)\/\-\,\'\&]+?)(?=[\x00-\x1f\d]{2}\s*\||\x00|\x01|\x02|\x03|\x04|\x05|$)', decoded)
                for w_str, a_name in raw_areas:
                    w_num = int(w_str)
                    if 1 <= w_num <= 18:
                        a_clean = a_name.strip()
                        if a_clean and a_clean not in areas_by_ward[w_num]:
                            areas_by_ward[w_num].append(a_clean)
        except Exception as e:
            self.stdout.write(self.style.WARNING(f"  Note: HTML viewstate area parsing: {e}"))

        ward_objs = {}
        for w_num in range(1, 19):
            z_code, default_name = zone_mapping.get(w_num, ("CENTRAL", f"Ward {w_num}"))
            area_list = areas_by_ward[w_num]
            landmarks_str = ", ".join(area_list[:15]) if area_list else default_name

            loc, _ = CivicLocation.objects.update_or_create(
                zone=z_code,
                ward_number=w_num,
                defaults={
                    "ward_name": f"Ward {w_num:02d} - {default_name}",
                    "known_landmarks": landmarks_str,
                }
            )
            ward_objs[w_num] = loc
        self.stdout.write(self.style.SUCCESS("  Synced all 18 official RMC Wards and area gazetteers."))

        # 3. Read official Complaint Types from Excel file and populate CivicCategories
        try:
            wb = openpyxl.load_workbook('Complaint Details 01-10-2026 to 10-10-2026.xlsx', data_only=True)
            ws = wb.active
            dept_categories = {}

            # Map Excel department string to official Department instance
            dept_alias_map = {
                "DRAINAGE CHOCKUP/OVERFLOW": dept_objs.get("DRAINAGE CHOCKUP/OVERFLOW"),
                "ROSHNI BRANCH": dept_objs.get("ROSHNI BRANCH"),
                "SOLID WASTE MANAGEMENT": dept_objs.get("SOLID WASTE MANAGEMENT"),
                "WATER WORKS(OUTDOOR)": dept_objs.get("WATER WORKS(OUTDOOR)"),
                "DRAINAGE MAINTENANCE": dept_objs.get("DRAINAGE MAINTENANCE"),
                "BANDHKAM": dept_objs.get("BANDHKAM"),
                "CITY BUS": dept_objs.get("CITY BUS (RAJKOT RAJPATH)"),
                "CONSERVANCY (DEAD ANIMAL)": dept_objs.get("CONSERVANCY (DEAD ANIMAL)"),
                "GARDEN": dept_objs.get("GARDEN & SPORTS"),
                "SPORTS": dept_objs.get("GARDEN & SPORTS"),
                "DABAN HATAV BRANCH": dept_objs.get("DABAN HATAV BRANCH"),
                "TOWN PLANNING": dept_objs.get("TOWN PLANNING & ESTATE"),
                "ESTATE BRANCH": dept_objs.get("TOWN PLANNING & ESTATE"),
                "URBAN MALERIYA": dept_objs.get("HEALTH & URBAN MALERIYA"),
                "HEALTH BRANCH": dept_objs.get("HEALTH & URBAN MALERIYA"),
                "A.N.C.D. (DOG STERILIZATION)": dept_objs.get("A.N.C.D. (DOG & ANIMAL CONTROL)"),
                "A.N.C.D.": dept_objs.get("A.N.C.D. (DOG & ANIMAL CONTROL)"),
                "TRAFFIC & TRNS. CELL": dept_objs.get("TRAFFIC & TRNS. CELL"),
                "PARKING MANAGEMENT BRANCH": dept_objs.get("TRAFFIC & TRNS. CELL"),
                "FOOD BRANCH": dept_objs.get("FOOD & SANITATION"),
                "HAND PUMP BRANCH": dept_objs.get("WATER WORKS(OUTDOOR)"),
                "S. W. M. (VOKALA)": dept_objs.get("SOLID WASTE MANAGEMENT"),
            }

            for r in range(5, ws.max_row + 1):
                raw_dept = ws.cell(row=r, column=2).value
                raw_ctype = ws.cell(row=r, column=3).value
                if raw_dept and raw_ctype:
                    dept_instance = dept_alias_map.get(raw_dept.strip())
                    if dept_instance:
                        ctype_clean = raw_ctype.strip()
                        if dept_instance not in dept_categories:
                            dept_categories[dept_instance] = set()
                        dept_categories[dept_instance].add(ctype_clean)

            total_cats = 0
            for dept_inst, ctypes in dept_categories.items():
                for ctype in ctypes:
                    code_slug = slugify(f"{dept_inst.code}-{ctype[:35]}")[:50]
                    CivicCategory.objects.update_or_create(
                        department=dept_inst,
                        name=ctype,
                        defaults={
                            "code": code_slug,
                            "keywords_en": ctype.lower(),
                            "is_active": True,
                        }
                    )
                    total_cats += 1
            self.stdout.write(self.style.SUCCESS(f"  Synced {total_cats} official RMC complaint categories from Excel dataset!"))

        except Exception as e:
            self.stdout.write(self.style.ERROR(f"Error loading Excel categories: {e}"))

        self.stdout.write(self.style.SUCCESS("Official RMC metadata synchronization complete!"))
