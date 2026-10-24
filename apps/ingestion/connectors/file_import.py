import os
import openpyxl
import csv
import hashlib
from datetime import datetime
from typing import Dict, Any, List, Optional, Tuple
from django.utils import timezone
from django.db import transaction

from .base import BaseConnector, NormalizedFeedbackPayload, IngestionBatchResult
from apps.feedback.models import FeedbackItem
from apps.ingestion.models import DataSource, ImportRun
from apps.intelligence.services import analyze_and_store_feedback

class FileImportConnector(BaseConnector):
    """
    Ingestion connector for batch spreadsheets (CSV and Excel).
    Specifically tailored for official RMC complaint exports and generic survey spreadsheets.
    Enforces strict data minimization (citizen names & mobile numbers are discarded).
    """
    connector_id = "file-import-connector"
    connector_name = "CSV / Excel File Importer"

    def validate_configuration(self, config: Dict[str, Any]) -> bool:
        return True

    def test_connection(self) -> Dict[str, Any]:
        return {"status": "HEALTHY", "message": "File import parser ready."}

    def normalize(self, raw_record: Dict[str, Any]) -> NormalizedFeedbackPayload:
        # Generic fallback normalization
        raw_text = raw_record.get("text", "") or raw_record.get("complaint", "")
        ext_id = str(raw_record.get("id", "") or hashlib.sha256(raw_text.encode()).hexdigest()[:12])
        return NormalizedFeedbackPayload(
            raw_text=raw_text,
            external_id=ext_id,
            source_type="FILE_IMPORT",
            published_at=timezone.now(),
            anonymized_sender_hash=hashlib.sha256(ext_id.encode()).hexdigest()[:16],
            is_synthetic=False
        )

    def process_file(
        self,
        file_path: str,
        data_source: DataSource,
        uploaded_by=None,
        max_rows: Optional[int] = None
    ) -> Tuple[ImportRun, IngestionBatchResult]:
        """
        Parses a CSV or Excel file, validates records, minimizes PII, deduplicates,
        stores FeedbackItem, and runs AI analysis.
        """
        result = IngestionBatchResult()
        file_name = os.path.basename(file_path)

        import_run = ImportRun.objects.create(
            data_source=data_source,
            uploaded_by=uploaded_by,
            file_name=file_name,
            status="PROCESSING"
        )

        try:
            if file_path.endswith((".xlsx", ".xls")):
                self._process_excel(file_path, data_source, import_run, result, max_rows)
            else:
                self._process_csv(file_path, data_source, import_run, result, max_rows)

            import_run.total_rows = result.total_processed
            import_run.imported_rows = result.new_items_count
            import_run.duplicate_rows = result.duplicates_blocked
            import_run.error_rows = result.errors_count
            import_run.status = "COMPLETED"
            import_run.error_log = result.error_details[:50]
            import_run.save()

            # Update data source last sync
            data_source.last_sync_at = timezone.now()
            data_source.save(update_fields=["last_sync_at"])

        except Exception as e:
            import_run.status = "FAILED"
            import_run.error_log = [f"Critical file processing error: {str(e)}"]
            import_run.save()
            raise

        return import_run, result

    def _process_excel(
        self,
        file_path: str,
        data_source: DataSource,
        import_run: ImportRun,
        result: IngestionBatchResult,
        max_rows: Optional[int] = None
    ):
        wb = openpyxl.load_workbook(file_path, data_only=True)
        ws = wb.active

        # Check if this is the official RMC complaint sheet (header on row 4)
        is_rmc_official = False
        start_row = 2
        for check_r in [1, 4]:
            row_vals = [str(ws.cell(row=check_r, column=c).value or "").strip() for c in range(1, 10)]
            if "Complaint No" in row_vals or "Complaint Type" in row_vals:
                is_rmc_official = True
                start_row = check_r + 1
                break

        for r in range(start_row, ws.max_row + 1):
            if max_rows and result.total_processed >= max_rows:
                break

            result.total_processed += 1
            try:
                if is_rmc_official:
                    payload = self._parse_official_rmc_row(ws, r)
                else:
                    payload = self._parse_generic_excel_row(ws, r)

                if not payload or not payload.raw_text.strip():
                    continue

                # Deduplication check via SHA-256 content_hash
                content_hash = payload.content_hash
                if FeedbackItem.objects.filter(content_hash=content_hash).exists():
                    result.duplicates_blocked += 1
                    continue

                # Save FeedbackItem
                with transaction.atomic():
                    item = FeedbackItem.objects.create(
                        data_source=data_source,
                        import_run=import_run,
                        external_id=payload.external_id,
                        content_hash=content_hash,
                        raw_text=payload.raw_text,
                        detected_language="und",
                        published_at=payload.published_at,
                        raw_metadata=payload.technical_metadata,
                        is_synthetic=payload.is_synthetic,
                        status="PENDING_ANALYSIS"
                    )
                    # Trigger immediate AI classification
                    analyze_and_store_feedback(item)

                result.new_items_count += 1

            except Exception as row_err:
                result.errors_count += 1
                result.error_details.append(f"Row {r}: {str(row_err)}")

    def _parse_official_rmc_row(self, ws, r: int) -> Optional[NormalizedFeedbackPayload]:
        """
        Parses official RMC complaint row:
        Col 2: Department, Col 3: Complaint Type, Col 4: Complaint No, Col 5: Complaint Date,
        Col 8: Zone, Col 9: Ward, Col 10: Address, Col 18: Remarks.
        STRICT DATA MINIMIZATION: Col 6 (Name) and Col 7 (Mobile No) are DISCARDED.
        """
        dept = ws.cell(row=r, column=2).value or ""
        ctype = ws.cell(row=r, column=3).value or ""
        cno = str(ws.cell(row=r, column=4).value or "")
        cdate_raw = str(ws.cell(row=r, column=5).value or "")
        zone = ws.cell(row=r, column=8).value or ""
        ward = str(ws.cell(row=r, column=9).value or "")
        address = ws.cell(row=r, column=10).value or ""
        remarks = ws.cell(row=r, column=18).value or ""

        if not ctype and not address:
            return None

        # Build clean civic text without citizen identity
        parts = []
        if dept:
            parts.append(f"[{dept}]")
        if ctype:
            parts.append(str(ctype))
        if address:
            parts.append(f"at {address}")
        if remarks and str(remarks).strip():
            parts.append(f"— Note: {remarks}")
        raw_text = " ".join(parts)

        # Parse date
        pub_date = timezone.now()
        if cdate_raw:
            for fmt in ["%d-%m-%Y %I:%M %p", "%d-%m-%Y %H:%M", "%Y-%m-%d %H:%M:%S", "%d/%m/%Y"]:
                try:
                    pub_date = timezone.make_aware(datetime.strptime(cdate_raw.strip(), fmt))
                    break
                except Exception:
                    pass

        # Anonymized hash
        sender_hash = hashlib.sha256(f"rmc-{cno}".encode()).hexdigest()[:16]

        return NormalizedFeedbackPayload(
            raw_text=raw_text,
            external_id=cno or f"rmc-row-{r}",
            source_type="FILE_IMPORT",
            published_at=pub_date,
            anonymized_sender_hash=sender_hash,
            technical_metadata={
                "official_rmc_export": True,
                "rmc_zone": str(zone),
                "rmc_ward": str(ward),
                "original_dept": str(dept),
                "original_complaint_type": str(ctype),
            },
            is_synthetic=False
        )

    def _parse_generic_excel_row(self, ws, r: int) -> Optional[NormalizedFeedbackPayload]:
        text = str(ws.cell(row=r, column=1).value or "").strip()
        if not text:
            return None
        return NormalizedFeedbackPayload(
            raw_text=text,
            external_id=f"row-{r}",
            source_type="FILE_IMPORT",
            published_at=timezone.now(),
            anonymized_sender_hash=hashlib.sha256(f"row-{r}".encode()).hexdigest()[:16],
            is_synthetic=False
        )

    def _process_csv(self, file_path, data_source, import_run, result, max_rows):
        with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
            reader = csv.DictReader(f)
            r = 1
            for row in reader:
                if max_rows and result.total_processed >= max_rows:
                    break
                result.total_processed += 1
                try:
                    text = row.get("text", "") or row.get("complaint", "") or list(row.values())[0]
                    if not text.strip():
                        continue
                    payload = NormalizedFeedbackPayload(
                        raw_text=text,
                        external_id=row.get("id", f"csv-{r}"),
                        source_type="FILE_IMPORT",
                        published_at=timezone.now(),
                        anonymized_sender_hash=hashlib.sha256(f"csv-{r}".encode()).hexdigest()[:16],
                        is_synthetic=False
                    )
                    content_hash = payload.content_hash
                    if FeedbackItem.objects.filter(content_hash=content_hash).exists():
                        result.duplicates_blocked += 1
                        continue

                    with transaction.atomic():
                        item = FeedbackItem.objects.create(
                            data_source=data_source,
                            import_run=import_run,
                            external_id=payload.external_id,
                            content_hash=content_hash,
                            raw_text=payload.raw_text,
                            detected_language="und",
                            published_at=payload.published_at,
                            raw_metadata=payload.technical_metadata,
                            is_synthetic=False,
                            status="PENDING_ANALYSIS"
                        )
                        analyze_and_store_feedback(item)

                    result.new_items_count += 1
                except Exception as e:
                    result.errors_count += 1
                    result.error_details.append(f"Row {r}: {str(e)}")
                r += 1
